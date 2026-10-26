"""Word report builders: resource-wise sprint report and product backlog report."""
import os
from collections import Counter
from datetime import date
from io import BytesIO

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

import docx_report_generator as g
from pie_chart import pie_png
from report_data import BUCKETS, backlog_items, by_resource, completion, load_items

STATUS_ORDER = {'Ready for Testing': 0, 'Ready for Retest': 0, 'In Progress': 1, 'New / Open': 2, 'HerdX Accepted': 3, 'Done': 4}
BUCKET_COLORS = {'Done': '0A963F', 'Testing': '1767C5', 'In Progress': 'E0A100', 'Open': '8FA3BD', 'Dropped': 'C9D2DD'}
COMPLETED_TYPES = [('User Stories', {'story', 'user story'}), ('Tasks', {'task', 'sub-task', 'subtask'}), ('Bugs', {'bug'})]


def _new_document(title, sprint, week):
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    for side in ('top_margin', 'bottom_margin', 'left_margin', 'right_margin'):
        setattr(section, side, Inches(0.625))
    section.header_distance = section.footer_distance = Inches(0.492)
    for part, text, align in ((section.header.paragraphs[0], f"Jira WSR | {sprint + ' ' if sprint else ''}{title}", WD_ALIGN_PARAGRAPH.LEFT),
                              (section.footer.paragraphs[0], f"{sprint} | Source: Jira CSV export | Page ", WD_ALIGN_PARAGRAPH.CENTER)):
        part.alignment = align
        run = part.add_run(text)
        run.font.name, run.font.size = 'Arial', Pt(6.5)
        run.font.color.rgb = RGBColor.from_string('9CA3AF')
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    section.footer.paragraphs[0]._p.append(field)
    normal = document.styles['Normal']
    normal.font.name, normal.font.size = 'Arial', Pt(8.5)
    normal.font.color.rgb = RGBColor.from_string(g.TEXT)

    heading = document.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    heading.paragraph_format.space_after = Pt(2)
    run = heading.add_run(f"{sprint + ' - ' if sprint else ''}{title}")
    run.bold, run.font.name, run.font.size = True, 'Arial', Pt(22)
    run.font.color.rgb = RGBColor.from_string(g.NAVY)
    sub = document.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.paragraph_format.space_after = Pt(10)
    run = sub.add_run(f"{week}   |   Prepared for stakeholder distribution")
    run.italic, run.font.name, run.font.size = True, 'Arial', Pt(11)
    run.font.color.rgb = RGBColor.from_string(g.MUTED)
    return document


def _note(document, text):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(8)
    run = paragraph.add_run(text)
    run.font.name, run.font.size = 'Arial', Pt(10)
    run.font.color.rgb = RGBColor.from_string(g.MUTED)


def _snapshot(document, tiles):
    table = _table(document, [[f"{value}\n{label}" for value, label, _ in tiles]], [g.CONTENT_WIDTH_IN / len(tiles)] * len(tiles), header=False, body_size=9)
    for cell, (_, _, color) in zip(table.rows[0].cells, tiles):
        lines = cell.text.split('\n')
        g._set_cell(cell, cell.text, align=WD_ALIGN_PARAGRAPH.CENTER, fill=g.LIGHT_BLUE, margin=110, size=9, runs=[(lines[0] + '\n', True, color), (lines[1], False, g.MUTED)])
    return table


def _table(document, rows, widths, **kw):
    table = g._table(document, rows, widths, **kw)
    for column, width in zip(table.columns, widths):  # LibreOffice/Google Docs read the grid, Word reads cell widths
        column.width = Inches(width)
    return table


def _completed_by_resource(document, groups, items):
    rows = [['Resource'] + [label for label, _ in COMPLETED_TYPES] + ['Total Done']]
    totals = Counter()
    for name, group in groups.items():
        done = [i for i in group if i['bucket'] == 'Done']
        counts = [sum(i['type'].lower() in types for i in done) for _, types in COMPLETED_TYPES]
        totals.update(dict(enumerate(counts + [len(done)])))
        rows.append([name] + [n or '–' for n in counts] + [str(len(done))])
    rows.append(['Total'] + [totals[k] for k in range(len(COMPLETED_TYPES) + 1)])
    first = 2.25
    _table(document, rows, [first] + [(g.CONTENT_WIDTH_IN - first) / 4] * 4, body_size=7.8)


def _completion_pie(document, items):
    counts = Counter(i['bucket'] for i in items)
    buckets = [b for b in BUCKETS if counts[b]]
    if not buckets:
        return
    layout = document.add_table(rows=1, cols=2)
    layout.autofit = False
    g._set_table_width(layout)
    chart_cell, legend_cell = layout.rows[0].cells
    for cell, width in ((chart_cell, 2.75), (legend_cell, g.CONTENT_WIDTH_IN - 2.75)):
        cell.width = Inches(width)
    for column, width in zip(layout.columns, (2.75, g.CONTENT_WIDTH_IN - 2.75)):
        column.width = Inches(width)
    picture = chart_cell.paragraphs[0]
    picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture.add_run().add_picture(BytesIO(pie_png([counts[b] for b in buckets], [BUCKET_COLORS[b] for b in buckets])), width=Inches(2.4))
    legend_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    rows = [['Status', 'Items', 'Share']] + [[b, counts[b], f'{round(100 * counts[b] / len(items))}%'] for b in buckets]
    legend = legend_cell.add_table(rows=len(rows), cols=3)
    legend.autofit = False
    g._borders(legend)
    widths = [2.1, 0.9, 0.9]
    g._set_table_width(legend, sum(widths))
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            cell = legend.cell(r, c)
            cell.width = Inches(widths[c])
            if r == 0:
                g._set_cell(cell, value, bold=True, color='FFFFFF', size=7.5, align=WD_ALIGN_PARAGRAPH.CENTER, fill=g.NAVY)
            elif c == 0:
                g._set_cell(cell, value, size=8.5, runs=[('■  ', True, BUCKET_COLORS[value]), (value, False, g.TEXT)])
            else:
                g._set_cell(cell, value, size=8.5, align=WD_ALIGN_PARAGRAPH.CENTER)
    for column, width in zip(legend.columns, widths):
        column.width = Inches(width)
    note = legend_cell.add_paragraph()
    note.paragraph_format.space_before = Pt(6)
    run = note.add_run(f'{completion(items)}% complete — {counts["Done"]} of {len(items) - counts["Dropped"]} active items done (Dropped excluded).')
    run.bold, run.font.name, run.font.size = True, 'Arial', Pt(9)
    run.font.color.rgb = RGBColor.from_string(g.GREEN)


def _sorted(items):
    return sorted(items, key=lambda i: (STATUS_ORDER.get(i['status'], 9), i['key']))


def _save(document, out):
    os.makedirs(os.path.dirname(out) or '.', exist_ok=True)
    document.save(out)


def build_resource_report(path, out, sprint='', week='Week not specified', filters=None):
    items = load_items(path, filters)
    groups = by_resource(items)
    assigned = [n for n in groups if n != 'Unassigned']
    document = _new_document('Resource-wise Sprint Report', sprint, week)

    g._heading(document, 'Sprint Snapshot')
    _snapshot(document, [(len(assigned), 'Resources', g.TEXT), (len(items), 'Work Items', g.BLUE),
                         (f'{completion(items)}%', 'Complete', g.GREEN), (len(groups.get('Unassigned', [])), 'Unassigned', g.RED)])

    g._heading(document, 'Workload by Resource')
    _note(document, 'Done includes HerdX Accepted. Testing covers Ready for Testing / Ready for Retest. Dropped covers Discarded, Rejected and Cannot Reproduce and is excluded from % Done.')
    rows = [['Resource', 'Total', 'Done', 'Testing', 'In Progress', 'Open', 'Dropped', '% Done']]
    totals = Counter()
    for name, group in groups.items():
        counts = Counter(i['bucket'] for i in group)
        totals.update(counts)
        rows.append([name, len(group)] + [counts[b] for b in BUCKETS] + [f'{completion(group)}%'])
    rows.append(['Total', len(items)] + [totals[b] for b in BUCKETS] + [f'{completion(items)}%'])
    widths = [1.9, 0.7, 0.7, 0.75, 0.9, 0.7, 0.8, 0.8]
    _table(document, rows, widths, body_size=7.8)

    g._heading(document, 'Completed by Resource')
    _note(document, 'User stories, tasks and bugs in Done or HerdX Accepted. Total Done also counts other work item types.')
    _completed_by_resource(document, groups, items)

    g._heading(document, 'Completion Breakdown')
    _completion_pie(document, items)

    types = [t for t, _ in Counter(i['type'] for i in items).most_common()]
    g._heading(document, 'Work Type Mix by Resource')
    rows = [['Resource'] + types]
    for name, group in groups.items():
        counts = Counter(i['type'] for i in group)
        rows.append([name] + [counts[t] or '–' for t in types])
    rows.append(['Total'] + [sum(1 for i in items if i['type'] == t) for t in types])
    first = 1.9
    _table(document, rows, [first] + [(g.CONTENT_WIDTH_IN - first) / len(types)] * len(types), body_size=7.8)

    g._heading(document, 'Resource Details')
    for name, group in groups.items():
        counts = Counter(i['bucket'] for i in group)
        p = document.add_paragraph()
        p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(10), Pt(3)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(f'{name}')
        run.bold, run.font.name, run.font.size = True, 'Arial', Pt(11)
        run.font.color.rgb = RGBColor.from_string(g.NAVY)
        run = p.add_run(f"   {len(group)} items · {counts['Done']} done · {counts['Testing']} testing · {counts['In Progress']} in progress · {counts['Open']} open")
        run.font.name, run.font.size = 'Arial', Pt(8)
        run.font.color.rgb = RGBColor.from_string(g.MUTED)
        rows = [['Key', 'Summary', 'Type', 'Status']] + [[i['key'], i['summary'], i['type'], i['status']] for i in _sorted(group)]
        _table(document, rows, [0.95, 3.75, 0.9, 1.65], status_column=3)
    _save(document, out)


def build_backlog_report(path, out, sprint='', week='Week not specified', filters=None):
    items = load_items(path, filters)
    backlog = [i for i in backlog_items(items) if i['type'].lower() == 'bug']
    today = date.today()
    ages = [(today - i['created']).days for i in backlog if i['created']]
    document = _new_document('Product Backlog Report', sprint, week)

    g._heading(document, 'Backlog Snapshot')
    _snapshot(document, [(len(backlog), 'Backlog Bugs', g.RED),
                         (sum(i['priority'].lower() in ('high', 'highest', 'critical', 'blocker') for i in backlog), 'High Priority', g.AMBER),
                         (sum(i['owner'] == 'Unassigned' for i in backlog), 'Unassigned', g.BLUE),
                         (f"{max(ages, default=0)}d", 'Oldest Bug', g.MUTED)])
    _note(document, f'Backlog = bugs not yet started (status New / Open). Other work item types are excluded. {len(backlog)} of {len(items)} items in this export qualify.')

    g._heading(document, 'Backlog by Priority')
    rows = [['Priority', 'Bugs', 'Unassigned', 'Oldest (days)']]
    for t, n in Counter(i['priority'] for i in backlog).most_common():
        group = [i for i in backlog if i['priority'] == t]
        rows.append([t, n, sum(i['owner'] == 'Unassigned' for i in group), max(((today - i['created']).days for i in group if i['created']), default=0)])
    _table(document, rows, [2.65, 1.5, 1.5, 1.6], body_size=8)

    g._heading(document, 'Backlog by Owner')
    rows = [['Owner', 'Backlog Bugs']]
    for name, group in by_resource(backlog).items():
        rows.append([name, len(group)])
    _table(document, rows, [4.6, 2.65], body_size=8)

    g._heading(document, f'Bug Backlog ({len(backlog)})')
    rows = [['Key', 'Summary', 'Owner', 'Priority', 'Created']]
    for i in sorted(backlog, key=lambda i: i['created'] or today):
        rows.append([i['key'], i['summary'], i['owner'], i['priority'], i['created'].strftime('%d %b %Y') if i['created'] else '—'])
    _table(document, rows, [0.95, 3.4, 1.2, 0.7, 1.0], body_size=7.4)

    g._heading(document, 'Next Focus')
    unassigned = sum(i['owner'] == 'Unassigned' for i in backlog)
    g._bullet(document, f'Assign owners to the {unassigned} unassigned backlog bugs.' if unassigned else 'All backlog bugs have an owner.')
    g._bullet(document, f"Review the {sum(1 for a in ages if a > 14)} bugs older than 14 days for priority or closure.")
    _save(document, out)
