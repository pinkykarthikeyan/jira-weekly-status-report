"""Weekly Status Report (Word): sprint snapshot plus QA and Dev resource detail tables."""
import re
from collections import Counter

from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.shared import Pt, RGBColor

import docx_report_generator as g
from report_data import by_resource, completion, load_items
from resource_report import _new_document, _save, _snapshot, _sorted, _table

DETAIL_WIDTHS = [0.95, 3.75, 0.9, 1.65]


def _name(value):
    return value.title() if value == value.lower() else value


def is_qa(group):
    """A resource counts as QA when any of their work items is a QA task."""
    return any(re.search(r'\bqa\b', i['summary'], re.I) for i in group)


def _resource_block(document, name, group):
    counts = Counter(i['bucket'] for i in group)
    p = document.add_paragraph()
    p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(10), Pt(3)
    p.paragraph_format.keep_with_next = True
    for text, bold, size in ((_name(name), True, 11),
                             (f"   {len(group)} items · {counts['Done']} done · {counts['Testing']} testing · {counts['In Progress']} in progress · {counts['Open']} open", False, 9)):
        run = p.add_run(text)
        run.bold, run.font.name, run.font.size = bold, 'Arial', Pt(size)
        run.font.color.rgb = RGBColor.from_string(g.TEXT)
        run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    rows = [['Key', 'Summary', 'Type', 'Status']] + [[i['key'], i['summary'], i['type'], i['status']] for i in _sorted(group)]
    table = _table(document, rows, DETAIL_WIDTHS, status_column=3, body_size=8)
    for row in table.rows[1:]:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER


def build_wsr_report(path, out, sprint='', week='Week not specified', filters=None):
    items = load_items(path, filters)
    groups = by_resource(items)
    unassigned = groups.pop('Unassigned', [])
    names = sorted(groups, key=str.lower)
    document = _new_document('Weekly Status Report', sprint, week)
    document.paragraphs[1].runs[0].text = week

    g._heading(document, 'Sprint Snapshot')
    _snapshot(document, [(len(groups), 'Resources', g.TEXT), (len(items), 'Work Items', g.BLUE),
                         (f'{completion(items)}%', 'Complete', g.GREEN), (len(unassigned), 'Unassigned', g.RED)])

    for title, members in (('QA Resource Details', [n for n in names if is_qa(groups[n])]),
                           ('Dev Resource Details', [n for n in names if not is_qa(groups[n])])):
        if not members:
            continue
        g._heading(document, title)
        for name in members:
            _resource_block(document, name, groups[name])
    if unassigned:
        g._heading(document, 'Unassigned Work Items')
        _resource_block(document, 'Unassigned', unassigned)
    _save(document, out)
